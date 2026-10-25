import base64
import io
import os
import uuid
from decimal import Decimal
import httpx
import pytest
from openpyxl import load_workbook
from app.api.imports import export_workbook
from app.features.xlsx_import import convert

HEADER = "sku,name,category_code,unit_code,barcode,price,cost,initial,reorder\n"


def test_editable_excel_roundtrip_keeps_literal_text_and_money():
    data = export_workbook(
        [
            dict(
                sku="ZERO-001",
                name="=1+1",
                category_code="general",
                unit_code="piece",
                barcode="001234",
                price=12345,
                cost=10001,
                reorder=Decimal("2.125"),
            )
        ]
    )
    book = load_workbook(io.BytesIO(data), data_only=False)
    assert book.active["B2"].data_type == "s" and book.active["B2"].value == "=1+1"
    assert book.active["E2"].value == "001234"
    assert book.active["H2"].value == 0
    result = convert(base64.b64encode(data).decode())
    assert "=1+1" in result["csv_text"] and "001234" in result["csv_text"]
    from app.api.imports import parse

    parsed = parse(result["csv_text"])[0][2]
    assert (
        parsed["price"] == 12345
        and parsed["cost"] == 10001
        and parsed["reorder"] == Decimal("2.125")
    )


@pytest.mark.skipif(
    not os.getenv("TEST_DATABASE_URL"), reason="Disposable PostgreSQL required"
)
@pytest.mark.asyncio
async def test_update_preview_conflicts_preservation_rollback_and_permissions():
    from app import config, db
    from app.main import app

    config.ADMIN_EMAIL = "test-admin@example.com"
    config.ADMIN_PASSWORD = "TestAdminPassword!2026"
    config.APP_ORIGIN = "http://testserver"
    db.DATABASE_URL = os.environ["TEST_DATABASE_URL"]
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url=config.APP_ORIGIN,
            headers={"Origin": config.APP_ORIGIN},
        ) as c:
            await c.post(
                "/api/auth/login",
                json=dict(email=config.ADMIN_EMAIL, password=config.ADMIN_PASSWORD),
            )

            async def old(action, **kw):
                r = await c.post("/api/workspace", json=dict(action=action, **kw))
                assert r.status_code == 200, r.text
                return r.json()

            tenant = (await old("tenant", name="Catalog update QA", currency="PKR"))[
                "id"
            ]
            base = f"/api/v1/tenants/{tenant}"
            url = base + "/catalog-imports"
            initial = (
                HEADER
                + "UP-1,Original,general,piece,00123,100,70,3,5\nUP-2,Second,general,piece,00456,200,90,2,1\n"
            )
            created = await c.post(
                url, json=dict(csv_text=initial, request=uuid.uuid4().hex)
            )
            assert created.status_code == 201, created.text
            first, second = created.json()["product_ids"]
            async with db.pool.acquire() as conn:
                await conn.execute(
                    "UPDATE products SET active=false,generic_name='Generic',strength='10mg',prescription_required=true,search_aliases='preserve me' WHERE id=$1",
                    first,
                )
                tax_id = uuid.uuid4().hex
                await conn.execute(
                    "INSERT INTO tax_profiles(id,tenant,name,rate_bps) VALUES($1,$2,'Retained tax',1200)",
                    tax_id,
                    tenant,
                )
                await conn.execute(
                    "UPDATE products SET tax_profile_id=$1 WHERE id=$2", tax_id, first
                )
                before_audit = await conn.fetchval(
                    "SELECT count(*) FROM audit WHERE tenant=$1", tenant
                )
            update = (
                HEADER
                + "up-1,Updated,general,piece,,125.50,71.25,0,6\nUP-2,Second updated,general,piece,00456,220,91,0,2\n"
            )
            preview = await c.post(
                url + "/preview", json=dict(csv_text=update, mode="update")
            )
            assert preview.status_code == 200, preview.text
            token = preview.json()["preview_token"]
            assert preview.json()["changes"][0]["changes"]["price"] == dict(
                before=10000, after=12550
            )
            async with db.pool.acquire() as conn:
                assert (
                    await conn.fetchval("SELECT price FROM products WHERE id=$1", first)
                    == 10000
                )
                assert (
                    await conn.fetchval(
                        "SELECT count(*) FROM audit WHERE tenant=$1", tenant
                    )
                    == before_audit
                )
            payload = dict(
                csv_text=update,
                mode="update",
                preview_token=token,
                request=uuid.uuid4().hex,
            )
            assert (
                await c.post(url, json={**payload, "preview_token": None})
            ).status_code == 409
            # Concurrent metadata edit invalidates the reviewed state and rolls back all rows.
            async with db.pool.acquire() as conn:
                await conn.execute(
                    "UPDATE products SET name='Concurrent edit' WHERE id=$1", second
                )
            stale = await c.post(url, json=payload)
            assert stale.status_code == 409, stale.text
            async with db.pool.acquire() as conn:
                assert (
                    await conn.fetchval("SELECT price FROM products WHERE id=$1", first)
                    == 10000
                )
                assert (
                    await conn.fetchval("SELECT name FROM products WHERE id=$1", second)
                    == "Concurrent edit"
                )
            payload["preview_token"] = (
                await c.post(
                    url + "/preview", json=dict(csv_text=update, mode="update")
                )
            ).json()["preview_token"]
            # Stock movement does not change a price/detail preview or get overwritten.
            await old(
                "movement",
                tenant=tenant,
                product=second,
                request=uuid.uuid4().hex,
                type="IN",
                quantity=2,
                notes="New stock",
            )
            posted = await c.post(url, json=payload)
            assert posted.status_code == 201, posted.text
            assert (await c.post(url, json=payload)).json() == posted.json()
            assert (
                await c.post(url, json={**payload, "mode": "create"})
            ).status_code == 409
            p = next(
                p
                for p in (await c.get(base + "/products")).json()["items"]
                if p["id"] == first
            )
            assert (
                p["stock"] == 3
                and p["price"] == 12550
                and p["cost"] == 7125
                and p["barcode"] is None
            )
            assert (
                p["active"] is False
                and p["generic_name"] == "Generic"
                and p["strength"] == "10mg"
                and p["prescription_required"]
                and p["search_aliases"] == "preserve me"
            )
            assert p["tax_profile_id"] == tax_id
            assert (
                next(
                    p
                    for p in (await c.get(base + "/products")).json()["items"]
                    if p["id"] == second
                )["stock"]
                == 4
            )
            exported = await c.get(url + "/export.xlsx", params=dict(q="UP-1"))
            assert (
                exported.status_code == 200
                and exported.headers["cache-control"] == "no-store"
            )
            wb = load_workbook(io.BytesIO(exported.content))
            assert wb.active["A2"].value == "UP-1" and wb.active.max_row == 2
            assert (
                await c.get(url + "/export.xlsx", params=dict(q="absent"))
            ).status_code == 404
            # Every failed file remains atomic, including a valid first row.
            for text in [
                update.replace("UP-2", "MISSING"),
                update.replace("UP-2", "UP-1"),
                update.replace("piece,,125.50,71.25,0", "piece,,125.50,71.25,1"),
                update.replace("piece,,125.50", "kg,,125.50"),
                update.replace("piece,,125.50", "piece,00456,125.50"),
            ]:
                r = await c.post(
                    url + "/preview", json=dict(csv_text=text, mode="update")
                )
                assert r.status_code in (400, 404), r.text
            other = (await old("tenant", name="Other catalog", currency="PKR"))["id"]
            assert (
                await c.post(
                    f"/api/v1/tenants/{other}/catalog-imports/preview",
                    json=dict(csv_text=update, mode="update"),
                )
            ).status_code == 404
            assert (
                await c.post(
                    url, json=payload, headers={"Origin": "http://evil.example"}
                )
            ).status_code == 403
            # Cashiers must not export/preview or commit catalog updates.
            email = uuid.uuid4().hex + "@example.com"
            await old(
                "member",
                tenant=tenant,
                email=email,
                name="Stock",
                password="StockPassword2026!",
                role="Cashier",
            )
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app),
                base_url=config.APP_ORIGIN,
                headers={"Origin": config.APP_ORIGIN},
            ) as staff:
                await staff.post(
                    "/api/auth/login",
                    json=dict(email=email, password="StockPassword2026!"),
                )
                assert (await staff.get(url + "/export.xlsx")).status_code == 403
                assert (
                    await staff.post(
                        url + "/preview", json=dict(csv_text=update, mode="update")
                    )
                ).status_code == 403
                assert (await staff.post(url, json=payload)).status_code == 403
            # Bound export, do not silently truncate a large catalog.
            async with db.pool.acquire() as conn:
                row = await conn.fetchrow("SELECT * FROM products WHERE id=$1", second)
                for n in range(199):
                    await conn.execute(
                        "INSERT INTO products(id,tenant,sku,name,category,price,category_id,unit_id) VALUES($1,$2,$3,'Bounded','General',1,$4,$5)",
                        uuid.uuid4().hex,
                        tenant,
                        f"BOUND-{n}",
                        row["category_id"],
                        row["unit_id"],
                    )
            assert (await c.get(url + "/export.xlsx")).status_code == 422
