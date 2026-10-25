"""Bounded XLSX-to-CSV decoding; all business validation stays in the CSV importer."""

import base64
import binascii
import csv
import re
from defusedxml.ElementTree import iterparse
from io import BytesIO, StringIO
from zipfile import ZipFile
from datetime import date, datetime, time
from fastapi import HTTPException
from openpyxl import load_workbook
from ..rules import fail

MAX_FILE_BYTES = 1_000_000
MAX_UNCOMPRESSED_BYTES = 8_000_000


def convert(encoded):
    from ..api.imports import COLUMNS, parse

    try:
        raw = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError):
        fail("Invalid XLSX file encoding.")
    if len(raw) > MAX_FILE_BYTES:
        fail("XLSX file is too large. Maximum 1 MB.")
    book = None
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 200
                or sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES
            ):
                fail("XLSX expanded content is too large.")
            if any(e.flag_bits & 1 for e in entries):
                fail("Encrypted workbooks are not supported.")
            if any(
                "vbaproject" in e.filename.lower()
                or "externallinks/" in e.filename.lower()
                for e in entries
            ):
                fail("Macros and external workbook links are not supported.")
            for entry in entries:
                if entry.filename.startswith(
                    "xl/worksheets/"
                ) and entry.filename.endswith(".xml"):
                    previous_row = 0
                    for event, element in iterparse(
                        BytesIO(archive.read(entry)), events=("end",)
                    ):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "c":
                            address = element.attrib.get("r", "")
                            match = re.fullmatch(r"[A-I]([1-9][0-9]{0,2})", address)
                            if not match or int(match.group(1)) > 201:
                                fail(
                                    "Use only template columns A–I and at most 200 product rows."
                                )
                        elif tag == "row":
                            row_number = int(element.attrib.get("r", "0"))
                            if not previous_row < row_number <= 201:
                                fail(
                                    "Use ordered worksheet rows and at most 200 products."
                                )
                            previous_row = row_number
                        element.clear()
        book = load_workbook(
            BytesIO(raw), read_only=True, data_only=False, keep_links=False
        )
        if len(book.worksheets) != 1:
            fail("Use one worksheet with the exact template columns.")
        sheet = book.worksheets[0]
        # Ignore unreliable declared dimensions; explicitly inspect a bounded region.
        sheet.reset_dimensions()
        output = StringIO()
        writer = csv.writer(output, lineterminator="\r\n")
        count = 0
        for number, cells in enumerate(
            sheet.iter_rows(min_row=1, max_row=201, max_col=9), start=1
        ):
            if number > 201:
                fail("Import at most 200 products at a time.")
            if len(cells) > 9 and any(c.value is not None for c in cells[9:]):
                fail("Use one worksheet with the exact template columns.")
            values = []
            for index, cell in enumerate(cells[:9]):
                value = cell.value
                if cell.data_type in {"f", "e"}:
                    fail(
                        f"XLSX row {number}: formulas and error cells are not supported; paste values instead."
                    )
                if isinstance(value, (datetime, date, time, bool)):
                    fail(
                        f"XLSX row {number}: use text identifiers and numeric amounts."
                    )
                if (
                    number > 1
                    and index in {0, 2, 3, 4}
                    and value is not None
                    and not isinstance(value, str)
                ):
                    fail(
                        f"XLSX row {number}: SKU, category, unit and barcode cells must be text to preserve identifiers."
                    )
                values.append("" if value is None else str(value))
            values += [""] * (9 - len(values))
            if number == 1:
                if values != COLUMNS:
                    fail("XLSX headers must match the downloadable template exactly.")
                writer.writerow(values)
            elif any(values):
                count += 1
                writer.writerow(values)
            if len(output.getvalue().encode("utf-8")) > 200_000:
                fail("Converted catalog is too large. Maximum 200 KB of CSV text.")
        result = output.getvalue()
        parse(result)  # Same decimal-money and product contract validation as CSV.
        return dict(csv_text=result, count=count)
    except HTTPException:
        raise
    except Exception:
        # Workbook metadata is untrusted. Never expose parser internals or paths.
        fail("Invalid or unsupported XLSX workbook.")
    finally:
        if book is not None:
            book.close()
