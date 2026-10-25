"""Application-runtime exports of authorized report data, with explicit units and bounds."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from math import ceil
import re
from ..rules import fail

# key, heading, kind. Money is converted from minor units only at presentation.
COLUMNS = {
    "products": [
        ("sku", "SKU", "text"),
        ("name", "Product", "text"),
        ("units", "Units", "number"),
        ("net_sales", "Net sales", "money"),
        ("cost", "Cost", "money"),
        ("profit", "Gross profit", "money"),
        ("margin_pct", "Margin (%)", "number"),
        ("abc", "ABC", "text"),
        ("xyz", "XYZ", "text"),
        ("cv", "CV", "number"),
        ("observed_days", "Observed days", "number"),
    ],
    "categories": [
        ("category", "Category", "text"),
        ("net_sales", "Net sales", "money"),
        ("cost", "Cost", "money"),
        ("profit", "Gross profit", "money"),
        ("units", "Units (mixed units)", "number"),
    ],
    "cashiers": [
        ("cashier", "Cashier", "text"),
        ("invoices", "Invoices", "number"),
        ("net_sales", "Net sales", "money"),
        ("cost", "Cost", "money"),
        ("profit", "Gross profit", "money"),
    ],
    "hours": [
        ("hour", "Hour", "number"),
        ("invoices", "Invoices", "number"),
        ("net_sales", "Net sales", "money"),
    ],
    "tax-summary": [
        ("profile", "Tax profile", "text"),
        ("rate_bps", "Rate (basis points)", "number"),
        ("inclusive", "Inclusive", "text"),
        ("taxable", "Net sales", "money"),
        ("tax", "Tax", "money"),
    ],
    "tax-invoices": [
        ("invoice", "Invoice", "text"),
        ("day", "Date", "text"),
        ("taxable", "Net sales", "money"),
        ("tax_included", "Tax included", "money"),
        ("tax_added", "Tax added", "money"),
        ("total", "Total", "money"),
    ],
    "stock-age": [
        ("sku", "SKU", "text"),
        ("name", "Product", "text"),
        ("lot", "Batch", "text"),
        ("units", "Quantity", "number"),
        ("unit", "Unit", "text"),
        ("age_days", "Age (days)", "number"),
        ("age_band", "Age band", "text"),
        ("expires", "Expiry", "text"),
        ("status", "Status", "text"),
        ("current_cost_value", "Current cost value", "money"),
    ],
    "basket-analysis": [
        ("sku_a", "SKU A", "text"),
        ("product_a", "Product A", "text"),
        ("sku_b", "SKU B", "text"),
        ("product_b", "Product B", "text"),
        ("baskets", "Shared baskets", "number"),
        ("support_pct", "Support (%)", "number"),
        ("confidence_a_b_pct", "Confidence A to B (%)", "number"),
        ("confidence_b_a_pct", "Confidence B to A (%)", "number"),
        ("lift", "Lift", "number"),
    ],
    "what-if": [
        ("baseline_net_sales", "Baseline net sales", "money"),
        ("baseline_profit", "Baseline gross profit", "money"),
        ("net_sales", "Projected net sales", "money"),
        ("cost", "Projected cost", "money"),
        ("profit", "Projected gross profit", "money"),
        ("profit_change", "Profit change", "money"),
        ("margin_pct", "Margin (%)", "number"),
    ],
    "turnover": [
        ("sku", "SKU", "text"),
        ("name", "Product", "text"),
        ("unit", "Unit", "text"),
        ("opening", "Opening stock", "number"),
        ("closing", "Closing stock", "number"),
        ("sold", "Sold units", "number"),
        ("returned", "Returned units", "number"),
        ("net_units", "Net sold units", "number"),
        ("average_stock", "Average daily closing stock", "number"),
        ("turnover", "Period turnover", "number"),
        ("days_per_turn", "Days per turn", "number"),
    ],
}
NOTES = {
    "management": "Completed invoices only; later returned invoices are excluded. Original sale-line costs; not FIFO valuation.",
    "tax": "Frozen taxes on completed invoices; later returned invoices excluded. Not a statutory filing or return-period ledger.",
    "stock-age": "Current stock at current catalog cost. Batch age begins at first creation, including replenished batches. Not FIFO/LIFO.",
    "basket-analysis": "Up to 200 qualifying pairs. Association is descriptive, not causation. Support/confidence values are percentages.",
    "what-if": "Independent demand, price and cost assumptions. No tax, overhead, stock constraints or price elasticity.",
    "turnover": "Net sold units / average daily closing stock, per SKU. Sales and returns use movement dates. Zero average produces blank turnover; non-positive turnover produces blank days. Not annualized or cost-based turnover.",
}


def prepare(name, data):
    if name in ("products", "categories", "cashiers", "hours"):
        rows = data[name]
        note = NOTES["management"]
    elif name.startswith("tax-"):
        rows = data["tax"]["groups" if name == "tax-summary" else "invoices"]
        note = NOTES["tax"]
    elif name == "stock-age":
        rows = data["stock"]["batches"]
        note = NOTES[name] + f" As of {data['stock']['as_of']}."
    elif name == "basket-analysis":
        rows = data["baskets"]["pairs"]
        note = NOTES[name] + f" Minimum baskets: {data['minimum_baskets']}."
    elif name == "what-if":
        rows = [data["scenario"]]
        note = (
            NOTES[name]
            + f" Horizon {data['scenario']['horizon_days']} days; demand {data['demand_pct']}%, price {data['price_pct']}%, cost {data['cost_pct']}%."
        )
    else:
        rows = data["items"]
        note = NOTES[name]
    return rows, COLUMNS[name], note


def value(raw, kind):
    if raw is None:
        return None
    if kind == "money":
        return Decimal(raw) / 100
    if kind == "number":
        return raw
    if isinstance(raw, bool):
        return "Yes" if raw else "No"
    return re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "\ufffd", str(raw)
    )


def xlsx_export(name, data):
    # Runs in the deployed Python application; no Codex authoring dependency.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    rows, columns, note = prepare(name, data)
    if len(rows) > 10000:
        fail("Excel export exceeds 10000 rows. Choose a shorter period.", 422)
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(
        [value(data["business"] + " | " + name.replace("-", " ").title(), "text")]
    )
    ws.append(
        [
            f"{data['start']} to {data['end']} | Asia/Karachi | {data['currency']} (major currency units)"
        ]
    )
    ws.append([note])
    for row in (1, 2, 3):
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=len(columns)
        )
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 32 if row < 3 else 48
        ws.cell(row, 1).data_type = "s"
    ws.cell(1, 1).font = Font(size=16, bold=True, color="17324D")
    ws.append(
        [c[1] + (f" ({data['currency']})" if c[2] == "money" else "") for c in columns]
    )
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 36
    for row_number, row in enumerate(rows, 5):
        ws.append([value(row.get(key), kind) for key, _, kind in columns])
        for column_number, (key, _, kind) in enumerate(columns, 1):
            cell = ws.cell(row_number, column_number)
            if kind == "text":
                cell.data_type = (
                    "s"  # identifiers and formula-like strings stay literal
                )
            else:
                integer_keys = {
                    "hour",
                    "invoices",
                    "baskets",
                    "rate_bps",
                    "age_days",
                    "observed_days",
                }
                numeric_format = (
                    "0"
                    if key in integer_keys
                    else {
                        "average_stock": "0.000000",
                        "cv": "0.0000",
                        "lift": "0.0000",
                        "turnover": "0.0000",
                        "days_per_turn": "0.00",
                    }.get(key, "0.00" if key.endswith("pct") else "0.000")
                )
                cell.number_format = (
                    "#,##0.00;[Red](#,##0.00)" if kind == "money" else numeric_format
                )
            cell.alignment = Alignment(vertical="top", wrap_text=kind == "text")
        ws.row_dimensions[row_number].height = max(
            18,
            15
            * max(
                (
                    ceil(len(str(row.get(k) or "")) / 24)
                    for k, _, kind in columns
                    if kind == "text"
                ),
                default=1,
            ),
        )
    for i, (_, _, kind) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 26 if kind == "text" else 20
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{max(4,ws.max_row)}"
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_export(name, data):
    from .. import pdf  # registers bundled Latin and Urdu fonts
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        LongTable,
        TableStyle,
    )
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib import colors
    from html import escape
    from arabic_reshaper import reshape
    from bidi.algorithm import get_display
    import re

    rows, columns, note = prepare(name, data)
    if len(rows) > 500:
        fail("PDF export exceeds 500 rows. Use Excel or a shorter period.", 422)
    styles = getSampleStyleSheet()
    for s in styles.byName.values():
        s.fontName = "StockDesk"
        s.fontSize = 9
        s.leading = 13
    styles["Title"].fontSize = 17
    styles["Title"].leading = 22
    styles["Title"].alignment = 0

    def p(v, style="BodyText"):
        text = value(v, "text") or ""
        if re.search("[\u0600-\u06ff]", text):
            text = get_display(reshape(text))
            parts = re.split("([\u0600-\u06ff\ufb50-\ufdff\ufe70-\ufeff]+)", text)
            text = "".join(
                (
                    '<font name="StockDeskUrdu">' + escape(x) + "</font>"
                    if re.search("[\u0600-\u06ff\ufb50-\ufdff\ufe70-\ufeff]", x)
                    else escape(x)
                )
                for x in parts
            )
        else:
            text = escape(text)
        return Paragraph(text, styles[style])

    def display(raw, kind):
        v = value(raw, kind)
        if v is None:
            return "-"
        if kind == "money":
            return f"{v:,.2f}"
        return str(v)

    out = BytesIO()
    page = landscape(A3)
    story = [
        p(data["business"] + " | " + name.replace("-", " ").title(), "Title"),
        p(
            f"{data['start']} to {data['end']} | Asia/Karachi | Amounts in {data['currency']}"
        ),
        p(note),
        Spacer(1, 12),
    ]
    grid = [[p(c[1]) for c in columns]] + [
        [p(display(r.get(k), kind)) for k, _, kind in columns] for r in rows
    ]
    if not rows:
        story.append(p("No matching records."))
    width = (page[0] - 60) / len(columns)
    table = LongTable(
        grid, colWidths=[width] * len(columns), repeatRows=1, hAlign="LEFT"
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EEF5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#DCE2EA")),
            ]
        )
    )
    story.append(table)

    def footer(canvas, doc):
        canvas.setFont("StockDesk", 8)
        canvas.drawRightString(page[0] - 30, 20, f"Page {doc.page}")

    SimpleDocTemplate(
        out, pagesize=page, leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=36
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
