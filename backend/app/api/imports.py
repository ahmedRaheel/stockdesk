"""Small atomic CSV imports. Preview exercises the real write rules in a rollback."""

import csv, io, json, hashlib
from decimal import Decimal, InvalidOperation
from fastapi import APIRouter, Depends, Request, HTTPException, Query, Response
from typing import Literal
from pydantic import Field, ValidationError
from .contracts import Contract, Identifier, ProductCreate
from .router import User, authenticated
from .. import db, service, security
from ..rules import fail
from ..features.catalog import save_product

router = APIRouter(
    prefix="/api/v1/tenants/{tenant}/catalog-imports",
    dependencies=[Depends(authenticated)],
    tags=["Catalog imports"],
)
COLUMNS = [
    "sku",
    "name",
    "category_code",
    "unit_code",
    "barcode",
    "price",
    "cost",
    "initial",
    "reorder",
]


class Preview(Contract):
    mode: Literal["create", "update"] = "create"
    csv_text: str = Field(min_length=1, max_length=200_000)


class Commit(Preview):
    request: Identifier
    preview_token: str | None = Field(default=None, max_length=64)


def parse(text):
    reader = csv.DictReader(io.StringIO(text.lstrip("\ufeff")), strict=True)
    rows = []
    try:
        if reader.fieldnames != COLUMNS:
            fail("CSV headers must match the downloadable template exactly.")
        for n, row in enumerate(reader, start=2):
            if n > 201:
                fail("Import at most 200 products at a time.")
            if None in row or any(v is None for v in row.values()):
                fail(f"CSV row {n}: incorrect column count.")
            row = {k: v.strip() for k, v in row.items()}
            try:
                # User-facing CSV money is major currency, never binary float.
                def money(key):
                    v = Decimal(row[key] or "0")
                    if (
                        not v.is_finite()
                        or v < 0
                        or v > 1_000_000_000
                        or v * 100 != (v * 100).to_integral_value()
                    ):
                        raise ValueError("Money requires at most two decimal places.")
                    return int(v * 100)

                payload = ProductCreate(
                    sku=row["sku"],
                    name=row["name"],
                    barcode=row["barcode"],
                    price=money("price"),
                    cost=money("cost"),
                    initial=row["initial"] or "0",
                    reorder=row["reorder"] or "0",
                )
            except (ValueError, InvalidOperation, ValidationError):
                fail(f"CSV row {n}: invalid product fields, money, or quantity.")
            rows.append((n, row, payload.model_dump()))
    except csv.Error:
        fail("Invalid CSV quoting or structure.")
    if not rows:
        fail("CSV contains no product rows.")
    return rows


async def apply(c, u, t, rows, mode="create"):
    changes, state, seen = [], [], set()
    ids = []
    for n, row, source in rows:
        payload = dict(source)
        try:
            sku = payload["sku"].upper()
            if sku in seen:
                fail("SKU appears more than once in this file.")
            seen.add(sku)
            current = None
            if mode == "update":
                if any(not row[k] for k in ("price", "cost", "reorder")):
                    fail(
                        "Update price, cost and reorder must be explicit; enter zero when intended."
                    )
                if payload["initial"] != 0:
                    fail(
                        "Opening stock must be zero for updates; use stock adjustments separately."
                    )
                current = await c.fetchrow(
                    "SELECT * FROM products WHERE tenant=$1 AND sku=$2", t, sku
                )
                if not current:
                    fail(
                        "Existing SKU not found; update imports never create products.",
                        404,
                    )
                state.append(
                    {
                        k: v
                        for k, v in dict(current).items()
                        if k not in ("stock", "stock_revision", "created")
                    }
                )

            for field, kind in [("category_code", "category"), ("unit_code", "unit")]:
                identifier = await c.fetchval(
                    "SELECT id FROM lookups WHERE tenant=$1 AND kind=$2 AND code=$3 AND active=true",
                    t,
                    kind,
                    row[field],
                )
                if not identifier:
                    fail(f"Unknown active {kind} code.")
                payload[kind + "_id"] = identifier
                if mode == "update":
                    state.append(
                        dict(
                            await c.fetchrow(
                                "SELECT * FROM lookups WHERE tenant=$1 AND id=$2",
                                t,
                                identifier,
                            )
                        )
                    )
            command = {**payload, "action": "product", "tenant": t}
            if current:
                # The file only owns these fields; retain tax, medical and lifecycle settings.
                editable = (
                    "name",
                    "barcode",
                    "price",
                    "cost",
                    "reorder",
                    "category_id",
                    "unit_id",
                )
                command = {
                    **dict(current),
                    **{k: payload[k] for k in editable},
                    "action": "editProduct",
                    "tenant": t,
                }
                diff = {
                    k: {"before": current[k], "after": command[k]}
                    for k in editable
                    if (current[k] or "") != (command[k] or "")
                }
                if "category_id" in diff:
                    diff["category_code"] = {
                        "before": await c.fetchval(
                            "SELECT code FROM lookups WHERE tenant=$1 AND id=$2",
                            t,
                            current["category_id"],
                        ),
                        "after": row["category_code"],
                    }
                    diff.pop("category_id")
                changes.append(
                    {"row": n, "sku": sku, "name": payload["name"], "changes": diff}
                )
            r = await save_product(c, u, command)
            ids.append(r["id"])
        except HTTPException as e:
            fail(f"CSV row {n}: {e.detail}", e.status_code)
    return ids, changes, state


def update_token(text, state):
    return hashlib.sha256(
        json.dumps({"csv": text, "state": state}, sort_keys=True, default=str).encode()
    ).hexdigest()


@router.post("/preview")
async def preview(tenant: str, payload: Preview, request: Request, user: User):
    security.csrf(request)
    rows = parse(payload.csv_text)
    async with db.pool.acquire() as c:
        async with c.transaction():
            await service.authorize(c, user, tenant, "product", lock=True)
            if payload.mode == "update":
                await service.authorize(c, user, tenant, "cost")
            # Explicit nested rollback leaves no product, movement or audit side effects.
            tx = c.transaction()
            await tx.start()
            try:
                _, changes, state = await apply(c, user, tenant, rows, payload.mode)
            finally:
                await tx.rollback()
    return {
        "count": len(rows),
        "mode": payload.mode,
        "changes": changes,
        "preview_token": (
            update_token(payload.csv_text, state) if payload.mode == "update" else None
        ),
        "items": [
            {
                "row": n,
                "sku": p["sku"].upper(),
                "name": p["name"],
                "price": p["price"],
                "initial": p["initial"],
            }
            for n, _, p in rows
        ],
    }


@router.post("", status_code=201)
async def commit(tenant: str, payload: Commit, request: Request, user: User):
    security.csrf(request)
    rows = parse(payload.csv_text)
    signature = hashlib.sha256(
        (
            payload.csv_text
            if payload.mode == "create"
            else "update:" + payload.csv_text
        ).encode()
    ).hexdigest()
    async with db.pool.acquire() as c:
        async with c.transaction():
            await service.authorize(c, user, tenant, "product", lock=True)
            if payload.mode == "update":
                await service.authorize(c, user, tenant, "cost")
            previous = await c.fetchrow(
                "SELECT * FROM catalog_imports WHERE tenant=$1 AND request=$2",
                tenant,
                payload.request,
            )
            if previous:
                if previous["signature"] != signature:
                    fail("Import request identifier already used.", 409)
                return {
                    "id": previous["id"],
                    "product_ids": json.loads(previous["product_ids"]),
                }
            ids, changes, state = await apply(c, user, tenant, rows, payload.mode)
            if payload.mode == "update" and payload.preview_token != update_token(
                payload.csv_text, state
            ):
                # Reject inside this transaction so all product/audit updates roll back.
                fail(
                    "Catalog changed or preview is missing. Validate and review the update again.",
                    409,
                )
            identifier = service.uid()
            await c.execute(
                "INSERT INTO catalog_imports(id,tenant,request,signature,product_ids,actor) VALUES($1,$2,$3,$4,$5,$6)",
                identifier,
                tenant,
                payload.request,
                signature,
                json.dumps(ids),
                user["email"],
            )
            await service.audit(
                c,
                tenant,
                user,
                (
                    "Catalog updated by import"
                    if payload.mode == "update"
                    else "Catalog imported"
                ),
                {
                    "id": identifier,
                    "count": len(ids),
                    "mode": payload.mode,
                    "changes": [
                        {
                            **x,
                            "changes": {
                                k: ("changed" if k == "cost" else v)
                                for k, v in x["changes"].items()
                            },
                        }
                        for x in changes
                    ],
                },
            )
    return {"id": identifier, "product_ids": ids}


class ExcelUpload(Contract):
    file_base64: str = Field(min_length=1, max_length=1_333_336)


@router.post("/convert-xlsx")
async def convert_excel(
    tenant: str, payload: ExcelUpload, request: Request, user: User
):
    import asyncio
    from ..features.xlsx_import import convert

    security.csrf(request)
    async with db.pool.acquire() as c:
        await service.authorize(c, user, tenant, "product")
    return await asyncio.to_thread(convert, payload.file_base64)


@router.get("/template.xlsx")
async def excel_template(tenant: str, user: User):
    from pathlib import Path
    from fastapi.responses import FileResponse

    async with db.pool.acquire() as c:
        await service.authorize(c, user, tenant, "product")
    return FileResponse(
        Path(__file__).resolve().parents[1] / "templates" / "catalog-template.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="catalog-template.xlsx",
    )


@router.get("/export.xlsx")
async def editable_export(tenant: str, user: User, q: str = Query("", max_length=100)):
    """Bounded editable snapshot. Text cell types prevent spreadsheet formula execution."""
    import asyncio

    async with db.pool.acquire() as c:
        async with c.transaction(isolation="repeatable_read", readonly=True):
            await service.authorize(c, user, tenant, "product")
            await service.authorize(c, user, tenant, "cost")
            records = service.rows(
                await c.fetch(
                    "SELECT p.*,cat.code AS category_code,u.code AS unit_code FROM products p JOIN lookups cat ON cat.tenant=p.tenant AND cat.id=p.category_id JOIN lookups u ON u.tenant=p.tenant AND u.id=p.unit_id WHERE p.tenant=$1 AND ($2='' OR strpos(lower(p.name||' '||p.sku||' '||COALESCE(p.barcode,'')),lower($2))>0) ORDER BY p.sku LIMIT 201",
                    tenant,
                    q,
                )
            )
    if not records:
        fail("No products match this export.", 404)
    if len(records) > 200:
        fail("More than 200 products match. Narrow the export search.", 422)
    content = await asyncio.to_thread(export_workbook, records)
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="catalog-update.xlsx"',
            "Cache-Control": "no-store",
        },
    )


def export_workbook(records):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    book = Workbook()
    sheet = book.active
    sheet.title = "Catalog"
    sheet.append(COLUMNS)
    for p in records:
        if any(
            ILLEGAL_CHARACTERS_RE.search(str(p[k] or ""))
            for k in ("sku", "name", "category_code", "unit_code", "barcode")
        ):
            fail(
                "Catalog text contains characters that Excel cannot store. Correct the product text before exporting.",
                422,
            )
        sheet.append(
            [
                p["sku"],
                p["name"],
                p["category_code"],
                p["unit_code"],
                p["barcode"] or "",
                Decimal(p["price"]) / 100,
                Decimal(p["cost"]) / 100,
                0,
                p["reorder"],
            ]
        )
        for cell in list(sheet[sheet.max_row])[:5]:
            cell.data_type = "s"
        for cell in list(sheet[sheet.max_row])[5:7]:
            cell.number_format = "0.00"
        for cell in list(sheet[sheet.max_row])[7:]:
            cell.number_format = "0.000"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24466C")
    for i, width in enumerate([22, 40, 22, 18, 25, 16, 16, 16, 16], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
